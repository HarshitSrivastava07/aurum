"""Strict, format-aware MCX file ingestion. Uploaded content is data, never code."""
from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import zipfile
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import PurePosixPath

import pandas as pd

MAX_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 250_000
MAX_MEMBERS = 50
SUPPORTED = {"GOLDM", "GOLDTEN", "GOLDGUINEA", "GOLDPETAL"}
NUMERIC = ["open", "high", "low", "close", "previous_close", "settlement", "volume", "open_interest", "trading_value"]
COLUMNS = ["symbol", "instrument", "date", "expiry", *NUMERIC, "source_file", "requested_date", "returned_date", "import_timestamp", "source_hash", "validation_status", "mode"]
ALIASES = {
    "symbol": ["symbol", "commodity", "ticker", "tckrsymb"],
    "instrument": ["instrument", "instrumenttype", "fininstrmtp", "instrumentname"],
    "date": ["date", "tradingdate", "tradedate", "timestamp", "tradDt", "bizDt"],
    "expiry": ["expiry", "expirydate", "expirationdate", "xpryDt", "fininstrmactlxprydt"],
    "open": ["open", "openprice", "opnpric"],
    "high": ["high", "highprice", "hghpric"],
    "low": ["low", "lowprice", "lwpric"],
    "close": ["close", "closeprice", "closingprice", "clspric"],
    "previous_close": ["previousclose", "prevclose", "prevcloseprice", "prvsclsgpric"],
    "settlement": ["settlement", "settlementprice", "settleprice", "sttlmpric"],
    "volume": ["volume", "vol", "vollots", "volumelots", "ttltradgvol"],
    "open_interest": ["openinterest", "oi", "oilots", "openinterestlots", "opnintrst"],
    "trading_value": ["tradingvalue", "value", "turnover", "ttltrfval", "ttltrdval", "totaltradedvalue"],
}


class ImportProblem(ValueError):
    def __init__(self, code: str, message: str, profile: dict | None = None):
        super().__init__(message)
        self.code, self.profile = code, profile or {}


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def clean_label(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def parse_date(value, order="mdy") -> str:
    """Parse declared numeric ordering, ISO, Excel datetimes, and 04SEP2026.

    Numeric text dates never silently switch ordering to rescue an invalid date.
    Excel serial numbers are deliberately not guessed as dates.
    """
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if order not in {"mdy", "dmy"}:
        raise ValueError("date_order must be mdy or dmy")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[ T]00:00:00)?", text):
        return datetime.strptime(text[:10], "%Y-%m-%d").date().isoformat()
    for fmt in ("%d%b%Y", "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    if re.fullmatch(r"\d{1,2}[/.-]\d{1,2}[/.-]\d{4}", text):
        sep = next(c for c in text if not c.isdigit())
        fmt = f"%m{sep}%d{sep}%Y" if order == "mdy" else f"%d{sep}%m{sep}%Y"
        return datetime.strptime(text, fmt).date().isoformat()
    raise ValueError(f"Unsupported or invalid date: {text!r}")


class _Tables(HTMLParser):
    """Read cell text only; do not execute scripts or fetch external resources."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self.rows, self.row, self.cell = [], None, None, None
        self.depth = 0
        self.ignored = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.ignored += 1
        if tag == "table":
            self.depth += 1
            if self.depth == 1:
                self.rows = []
        if self.depth == 1 and tag == "tr":
            self.row = []
        if self.depth == 1 and tag in {"td", "th"}:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None and not self.ignored:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.ignored = max(0, self.ignored - 1)
        if self.depth == 1 and tag in {"td", "th"} and self.cell is not None:
            if self.row is not None:
                self.row.append(" ".join(self.cell).strip())
            self.cell = None
        if self.depth == 1 and tag == "tr" and self.row is not None:
            if len(self.rows) >= MAX_ROWS:
                raise ImportProblem("too_many_rows", "File exceeds the row limit")
            self.rows.append(self.row)
            self.row = None
        if tag == "table":
            if self.depth == 1 and self.rows:
                self.tables.append(self.rows)
            self.depth -= 1


def _matrix_frame(rows) -> pd.DataFrame:
    """Find an explicit table header in the first 30 rows, never infer dates."""
    rows = list(rows)
    if len(rows) > MAX_ROWS:
        raise ImportProblem("too_many_rows", "File exceeds the row limit")
    for i, row in enumerate(rows[:30]):
        labels = {clean_label(x) for x in row if x is not None}
        if labels & {"symbol", "commodity", "tckrsymb"}:
            header = [str(x).strip() if x is not None else f"unnamed_{j}" for j, x in enumerate(row)]
            values = [list(r)[:len(header)] + [None] * max(0, len(header) - len(r)) for r in rows[i + 1:] if any(x not in (None, "") for x in r)]
            return pd.DataFrame(values, columns=header)
    raise ImportProblem("missing_header", "No supported Commodity/Symbol table header found")


def read_table(filename: str, content: bytes) -> tuple[pd.DataFrame, str]:
    if not content or len(content) > MAX_BYTES:
        raise ImportProblem("file_size", "File must contain 1 byte to 25 MiB")
    ext = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if ext not in {".csv", ".tsv", ".xls", ".xlsx", ".html", ".htm"}:
        raise ImportProblem("unsupported_format", "Supported formats: CSV, XLS, XLSX, HTML tables, ZIP")
    sample = content[:1024].lstrip().lower()
    if b"<table" in sample or b"<html" in sample or b"<!doctype html" in sample:
        parser = _Tables()
        parser.feed(content.decode("utf-8-sig", errors="replace"))
        for rows in parser.tables:
            try:
                return _matrix_frame(rows), "html-table"
            except ImportProblem as exc:
                if exc.code != "missing_header":
                    raise
        raise ImportProblem("missing_header", "No supported table in HTML export")
    if content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        import xlrd
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows > MAX_ROWS:
                    raise ImportProblem("too_many_rows", "File exceeds the row limit")
                rows = []
                for r in range(sheet.nrows):
                    rows.append([xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value for c in sheet.row(r)])
                try:
                    return _matrix_frame(rows), "xls-biff"
                except ImportProblem as exc:
                    if exc.code != "missing_header":
                        raise
        finally:
            book.release_resources()
        raise ImportProblem("missing_header", "No supported table in XLS workbook")
    if content.startswith(b"PK\x03\x04") and ext == ".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ImportProblem("zip_size", "Expanded workbook exceeds 100 MiB")
            if any(i.filename.endswith("vbaProject.bin") for i in archive.infolist()):
                raise ImportProblem("macros", "Macro-enabled workbooks are not supported")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book.worksheets:
                if (sheet.max_row or 0) > MAX_ROWS:
                    raise ImportProblem("too_many_rows", "File exceeds the row limit")
                try:
                    return _matrix_frame(sheet.iter_rows(values_only=True)), "xlsx"
                except ImportProblem as exc:
                    if exc.code != "missing_header":
                        raise
        finally:
            book.close()
        raise ImportProblem("missing_header", "No supported table in XLSX workbook")
    if ext in {".xlsx", ".xls"}:
        raise ImportProblem("file_signature", "File extension does not match a recognized Excel or HTML export")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")
    if "\x00" in text:
        raise ImportProblem("file_signature", "Binary content is not a CSV table")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    return _matrix_frame(csv.reader(io.StringIO(text), dialect)), "csv"


def zip_members(content: bytes) -> list[tuple[str, bytes]]:
    if len(content) > MAX_BYTES:
        raise ImportProblem("file_size", "Compressed archive exceeds 25 MiB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = [x for x in archive.infolist() if not x.is_dir()]
            if not entries or len(entries) > MAX_MEMBERS:
                raise ImportProblem("zip_members", "ZIP must contain between 1 and 50 files")
            if sum(x.file_size for x in entries) > MAX_EXPANDED_BYTES:
                raise ImportProblem("zip_size", "Expanded ZIP exceeds 100 MiB")
            for entry in entries:
                path = PurePosixPath(entry.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts or ":" in entry.filename:
                    raise ImportProblem("zip_path", "ZIP contains an unsafe path")
                if entry.flag_bits & 1:
                    raise ImportProblem("zip_encrypted", "Encrypted ZIP files are not supported")
                if path.suffix.lower() not in {".csv", ".tsv", ".xls", ".xlsx", ".html", ".htm"}:
                    raise ImportProblem("zip_member_type", "ZIP contains unsupported or nested archive files")
                if entry.file_size > MAX_BYTES or entry.file_size / max(entry.compress_size, 1) > 250:
                    raise ImportProblem("zip_ratio", "ZIP member exceeds size or compression-ratio limit")
                if (entry.external_attr >> 16) & 0o170000 == 0o120000:
                    raise ImportProblem("zip_symlink", "ZIP symbolic links are not supported")
            return [(x.filename, archive.read(x)) for x in entries]
    except zipfile.BadZipFile as exc:
        raise ImportProblem("invalid_zip", "Invalid ZIP archive") from exc


def normalize_table(frame: pd.DataFrame, filename: str, content: bytes, requested_date=None,
                    date_order="mdy", mode="real", today: date | None = None) -> tuple[pd.DataFrame, dict]:
    if mode not in {"real", "synthetic"}:
        raise ImportProblem("invalid_mode", "mode must be real or synthetic")
    if date_order not in {"mdy", "dmy"}:
        raise ImportProblem("invalid_date_order", "date_order must be mdy or dmy")
    today = today or datetime.now().date()
    req = None
    if requested_date:
        try:
            req = parse_date(requested_date, "dmy")
        except ValueError as exc:
            raise ImportProblem("invalid_requested_date", str(exc)) from exc
        if req > today.isoformat():
            raise ImportProblem("future_requested_date", "Requested date is in the future")
    labels = {clean_label(c): c for c in frame.columns}
    profile = {"columns": [str(c) for c in frame.columns], "source_rows": len(frame)}
    if "year" in labels and "date" not in labels and "tradDt".lower() not in labels:
        profile["years"] = sorted({str(v).strip() for v in frame[labels["year"]].dropna()})
        profile["kind"] = "annual_aggregate"
        raise ImportProblem("annual_aggregate", "This is a YearWise annual turnover report, not daily Bhavcopy. It has no session date, expiry or daily prices and cannot support a historical strategy.", profile)
    mapping = {}
    for canonical, aliases in ALIASES.items():
        found = [labels[clean_label(a)] for a in aliases if clean_label(a) in labels]
        if found:
            mapping[canonical] = found[0]
    missing = [x for x in ("symbol", "date", "expiry", "close") if x not in mapping]
    if missing:
        raise ImportProblem("missing_columns", "Missing required daily fields: " + ", ".join(missing), profile)
    digest, now = sha256(content), datetime.now(timezone.utc).isoformat()
    output, issues, rejected, skipped = [], [], 0, 0
    for idx, original in frame.iterrows():
        rownum = int(idx) + 2
        raw = {k: original[v] for k, v in mapping.items()}
        symbol = str(raw.get("symbol", "")).strip().upper()
        if symbol not in SUPPORTED:
            skipped += 1
            continue
        instrument = str(raw.get("instrument", "FUTCOM")).strip().upper()
        if instrument not in {"FUTCOM", "FUTCOMM", "FUT", "FUTURES", "FUTCOMMODITY"}:
            skipped += 1
            continue
        errors, warnings = [], []
        row = {c: None for c in COLUMNS}
        row.update(symbol=symbol, instrument="FUTCOM", source_file=filename, requested_date=req,
                   source_hash=digest, import_timestamp=now, mode=mode)
        for field in ("date", "expiry"):
            try:
                row[field] = parse_date(raw[field], date_order)
            except (ValueError, TypeError):
                errors.append(f"invalid_{field}")
        row["returned_date"] = row["date"]
        if row["date"] and row["date"] > today.isoformat():
            errors.append("future_observation")
        if row["date"] and row["expiry"] and row["date"] > row["expiry"]:
            errors.append("observation_after_expiry")
        if req and row["date"] and req != row["date"]:
            warnings.append("requested_returned_date_mismatch")
        for field in NUMERIC:
            val = raw.get(field)
            if val is None or pd.isna(val) or str(val).strip() in {"", "-", "--", "NA", "N/A", "null"}:
                if field == "close":
                    errors.append("missing_close")
                continue
            try:
                number = float(str(val).strip().replace(",", ""))
                if not math.isfinite(number):
                    raise ValueError()
                if field in {"volume", "open_interest", "trading_value"}:
                    if number < 0:
                        raise ValueError()
                    if field in {"volume", "open_interest"} and number != int(number):
                        raise ValueError()
                elif number <= 0:
                    raise ValueError()
                row[field] = number
            except (ValueError, TypeError):
                errors.append(f"invalid_{field}")
        high, low = row["high"], row["low"]
        if high is not None and low is not None and high < low:
            errors.append("high_below_low")
        for field in ("open", "close"):
            if row[field] is not None and ((high is not None and row[field] > high) or (low is not None and row[field] < low)):
                errors.append(f"{field}_outside_range")
        if row["open"] is None:
            warnings.append("missing_open_execution_unavailable")
        if row["volume"] is None or row["open_interest"] is None:
            warnings.append("missing_liquidity")
        if row["settlement"] is None:
            warnings.append("settlement_not_supplied_close_not_substituted")
        if errors:
            rejected += 1
            issues.append({"row": rownum, "symbol": symbol, "severity": "error", "codes": sorted(set(errors))})
        else:
            row["validation_status"] = "warning" if warnings else "valid"
            output.append(row)
            if warnings:
                issues.append({"row": rownum, "symbol": symbol, "severity": "warning", "codes": warnings})
    result = pd.DataFrame(output, columns=COLUMNS)
    for field in NUMERIC:
        result[field] = pd.to_numeric(result[field], errors="coerce").astype(float)
    return result, {"source_rows": len(frame), "rows_rejected": rejected, "rows_skipped": skipped,
                    "issues": issues[:500], "issue_count": len(issues), "profile": profile,
                    "date_order": date_order, "requested_date": req}
